"""Real-document regressions for DOCX hyperlink display text extraction."""

from pathlib import Path
from typing import cast

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml.etree import SubElement, _Element

from gptme.tools.convert import convert_file


def _add_link(para: Paragraph, *labels: tuple[str, bool, bool]) -> None:
    """Add an internal hyperlink without fetching or relying on a remote target."""
    link = cast(_Element, SubElement(para._p, qn("w:hyperlink")))
    link.set(qn("w:anchor"), "bookmark")
    for text, bold, italic in labels:
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        if bold:
            properties.append(OxmlElement("w:b"))
        if italic:
            properties.append(OxmlElement("w:i"))
        run.append(properties)
        value = OxmlElement("w:t")
        value.text = text
        run.append(value)
        link.append(run)


@pytest.mark.parametrize(
    ("style", "prefix"),
    [
        (None, ""),
        ("Heading 1", "# "),
        ("Heading 2", "## "),
        ("Heading 3", "### "),
        ("List Bullet", "- "),
    ],
)
def test_docx_markdown_keeps_link_labels_in_inline_order(
    tmp_path: Path, style: str | None, prefix: str
) -> None:
    doc = Document()
    para = doc.add_paragraph("before ", style=style)
    _add_link(para, ("FIRST", False, False))
    para.add_run(" between ")
    _add_link(para, ("SECOND", False, False))
    para.add_run(" after")
    source = tmp_path / "links.docx"
    doc.save(str(source))
    destination = tmp_path / "links.md"

    result = convert_file(source, destination)

    assert result.success, result.error
    assert result.converter_used == "python-docx"
    assert destination.read_text() == f"{prefix}before FIRST between SECOND after"


def test_docx_markdown_link_only_and_empty_runs(tmp_path: Path) -> None:
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("")
    _add_link(para, ("", False, False), ("ONLY LINK", False, False))
    para.add_run("")
    doc.add_paragraph()
    _add_link(doc.add_paragraph(), ("SECOND LINK", False, False))
    source = tmp_path / "only.docx"
    doc.save(str(source))
    destination = tmp_path / "only.md"

    result = convert_file(source, destination)

    assert result.success, result.error
    assert destination.read_text() == "ONLY LINK\n\nSECOND LINK"


def test_docx_markdown_link_run_formatting_and_plain_control(tmp_path: Path) -> None:
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("plain ")
    para.add_run("bold").bold = True
    para.add_run(" ")
    _add_link(
        para,
        ("linked bold", True, False),
        (" ", False, False),
        ("linked italic", False, True),
    )
    para.add_run(" ")
    para.add_run("italic").italic = True
    doc.add_paragraph("ordinary paragraph")
    source = tmp_path / "formatted.docx"
    doc.save(str(source))
    destination = tmp_path / "formatted.md"

    result = convert_file(source, destination)

    assert result.success, result.error
    assert (
        destination.read_text()
        == "plain **bold** **linked bold** *linked italic* *italic*\nordinary paragraph"
    )


def test_docx_txt_unchanged_with_links(tmp_path: Path) -> None:
    doc = Document()
    para = doc.add_paragraph("ordinary paragraph ")
    _add_link(para, ("LINK", False, False))
    para.add_run(" after")
    source = tmp_path / "plain.docx"
    doc.save(str(source))
    destination = tmp_path / "plain.txt"

    result = convert_file(source, destination)

    assert result.success, result.error
    # TXT deliberately keeps python-docx's existing text policy, including
    # version-dependent hyperlink support; this repair changes Markdown only.
    assert destination.read_text() == Document(str(source)).paragraphs[0].text
