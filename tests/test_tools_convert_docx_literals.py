"""Rendered semantics of literal Word paragraphs, using saved/reloaded DOCX."""

from pathlib import Path
from types import SimpleNamespace

import pytest
from markdown_it import MarkdownIt

pytest.importorskip("docx")
from docx import Document

from gptme.tools.convert import _para_to_md, convert_file


@pytest.mark.parametrize(
    "text",
    [
        "# literal heading",
        "- literal list",
        "1. literal numbered list",
        "[literal](https://example.invalid)",
        "<b>literal HTML</b>",
        "*literal emphasis* and **bold syntax**",
        "`literal code` and \\backslash",
        "&copy; stays literal",
        "---",
        "    not a code block",
    ],
)
def test_docx_literal_text_is_not_markdown(tmp_path: Path, text: str) -> None:
    document = Document()
    para = document.add_paragraph()
    # Every character in a separate run catches syntax split across run edges.
    for char in text:
        para.add_run(char)
    source = tmp_path / "literal.docx"
    document.save(str(source))
    output = tmp_path / "literal.md"
    result = convert_file(source, output)
    assert result.success, result.error
    tokens = MarkdownIt("commonmark").parse(output.read_text())
    assert [token.type for token in tokens] == [
        "paragraph_open",
        "inline",
        "paragraph_close",
    ]
    children = tokens[1].children or []
    assert all(child.type == "text" for child in children)
    assert "".join(child.content for child in children) == text


@pytest.mark.parametrize("indent", ["    ", "\t", " \t  "])
def test_manual_break_indentation_stays_literal(tmp_path: Path, indent: str) -> None:
    document = Document()
    para = document.add_paragraph("intro")
    para.add_run().add_break()
    para.add_run().add_break()
    # The indentation and literal syntax can span multiple runs.
    for char in indent + "*literal*":
        para.add_run(char)
    source = tmp_path / "breaks.docx"
    document.save(str(source))
    md = tmp_path / "breaks.md"
    txt = tmp_path / "breaks.txt"
    assert convert_file(source, md).success
    assert convert_file(source, txt).success
    assert txt.read_text() == "intro\n\n" + indent + "*literal*"
    tokens = MarkdownIt("commonmark").parse(md.read_text())
    assert [token.type for token in tokens] == [
        "paragraph_open",
        "inline",
        "paragraph_close",
        "paragraph_open",
        "inline",
        "paragraph_close",
    ]
    children = tokens[4].children or []
    assert all(child.type == "text" for child in children)
    assert "".join(child.content for child in children) == indent + "*literal*"


@pytest.mark.parametrize("style", ["List Bullet", "List Number"])
@pytest.mark.parametrize("indent", ["", "    ", "\t", " \t  "])
@pytest.mark.parametrize("breaks", [1, 2])
def test_list_manual_breaks_stay_inside_item(
    tmp_path: Path, style: str, indent: str, breaks: int
) -> None:
    document = Document()
    para = document.add_paragraph("intro", style=style)
    for _ in range(breaks):
        para.add_run().add_break()
    for char in indent + "*literal*":
        para.add_run(char)
    document.add_paragraph("outside")
    source = tmp_path / "list-breaks.docx"
    document.save(str(source))
    md = tmp_path / "list-breaks.md"
    txt = tmp_path / "list-breaks.txt"
    assert convert_file(source, md).success
    assert convert_file(source, txt).success
    assert txt.read_text() == "intro" + "\n" * breaks + indent + "*literal*\noutside"

    tokens = MarkdownIt("commonmark").parse(md.read_text())
    assert sum(t.type == "list_item_open" for t in tokens) == 1
    assert not any(t.type in ("code_block", "fence") for t in tokens)
    item_start = next(i for i, t in enumerate(tokens) if t.type == "list_item_open")
    item_end = next(i for i, t in enumerate(tokens) if t.type == "list_item_close")
    item_children = [c for t in tokens[item_start:item_end] for c in (t.children or [])]
    assert all(c.type in ("text", "softbreak") for c in item_children)
    assert [c.content for c in item_children if c.type == "text"] == [
        "intro",
        indent + "*literal*",
    ]
    assert [t.content for t in tokens[item_end:] if t.type == "inline"] == ["outside"]


def test_separate_paragraphs_and_unchanged_txt(tmp_path: Path) -> None:
    document = Document()
    paragraphs = ["first paragraph", "second paragraph", "", "third paragraph"]
    for text in paragraphs:
        document.add_paragraph(text)
    source = tmp_path / "paragraphs.docx"
    document.save(str(source))
    md = tmp_path / "paragraphs.md"
    txt = tmp_path / "paragraphs.txt"
    assert convert_file(source, md).success
    assert convert_file(source, txt).success
    assert txt.read_text() == "\n".join(paragraphs)
    tokens = MarkdownIt().parse(md.read_text())
    assert [t.content for t in tokens if t.type == "inline"] == [
        text for text in paragraphs if text
    ]
    assert sum(t.type == "paragraph_open" for t in tokens) == 3


def test_intentional_word_styles_survive_literal_escaping(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("# literal", style="Heading 1")
    document.add_paragraph("- literal", style="List Bullet")
    para = document.add_paragraph()
    para.add_run("*bold literal*").bold = True
    para.add_run(" ")
    para.add_run("_italic literal_").italic = True
    source = tmp_path / "styles.docx"
    document.save(str(source))
    output = tmp_path / "styles.md"
    assert convert_file(source, output).success
    tokens = MarkdownIt().parse(output.read_text())
    assert any(t.type == "heading_open" for t in tokens)
    assert any(t.type == "bullet_list_open" for t in tokens)
    children = [c for t in tokens for c in (t.children or [])]
    assert sum(c.type == "strong_open" for c in children) == 1
    assert sum(c.type == "em_open" for c in children) == 1
    texts = [c.content for c in children if c.type == "text"]
    assert "# literal" in texts and "- literal" in texts
    assert "*bold literal*" in texts and "_italic literal_" in texts


def test_fallback_text_and_blank_paragraph() -> None:
    para = SimpleNamespace(text="[literal] &copy;", runs=[], style=None)
    tokens = MarkdownIt().parse(_para_to_md(para))
    assert "".join(c.content for c in tokens[1].children or []) == para.text
    assert all(c.type == "text" for c in tokens[1].children or [])
    assert _para_to_md(SimpleNamespace(text="  ", runs=[], style=None)) == ""
