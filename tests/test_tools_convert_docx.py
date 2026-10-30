"""Real DOCX fixtures for document text extraction (optional python-docx)."""

from __future__ import annotations

from typing import TYPE_CHECKING
from unittest.mock import patch

import pytest

docx = pytest.importorskip("docx")

from gptme.tools.convert import DocumentToTextConverter, ToolAvailability

if TYPE_CHECKING:
    from pathlib import Path

    from docx.document import Document


def convert(document: Document, tmp_path: Path, extension: str) -> str:
    source = tmp_path / "document.docx"
    target = tmp_path / f"document.{extension}"
    document.save(str(source))
    with patch(
        "gptme.tools.convert.get_availability",
        return_value=ToolAvailability(python_docx=True, libreoffice=False),
    ):
        result = DocumentToTextConverter().convert(source, target)
    assert result.success, result.error
    assert result.converter_used == "python-docx"
    return target.read_text()


@pytest.mark.parametrize("extension", ["txt", "md"])
@pytest.mark.parametrize("interleaved", [False, True])
def test_docx_table_text_in_body_order(
    tmp_path: Path, extension: str, interleaved: bool
) -> None:
    document = docx.Document()
    if interleaved:
        document.add_paragraph("BEFORE_TABLE")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("ITEM", "AMOUNT"), ("INVOICE", "123.45")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    if interleaved:
        document.add_paragraph("AFTER_TABLE")
    text = convert(document, tmp_path, extension)
    expected = (
        "ITEM\tAMOUNT\nINVOICE\t123.45"
        if extension == "txt"
        else "| ITEM | AMOUNT |\n| --- | --- |\n| INVOICE | 123.45 |"
    )
    if interleaved:
        separator = "\n\n" if extension == "md" else "\n"
        expected = separator.join(["BEFORE_TABLE", expected, "AFTER_TABLE"])
    assert text == expected


@pytest.mark.parametrize("extension", ["txt", "md"])
@pytest.mark.parametrize(
    ("short_row", "before", "after"), [(0, 1, 0), (0, 0, 1), (1, 1, 1)]
)
def test_docx_omitted_grid_positions(
    tmp_path: Path, extension: str, short_row: int, before: int, after: int
) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = docx.Document()
    table = document.add_table(rows=2, cols=3)
    expected_rows = [[f"R{r}C{c}" for c in range(3)] for r in range(2)]
    for row, values in zip(table.rows, expected_rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
    row = table.rows[short_row]
    # Omitted positions are absent cells, not existing cells with empty text.
    for name, count in [("gridBefore", before), ("gridAfter", after)]:
        if count:
            prop = OxmlElement(f"w:{name}")
            prop.set(qn("w:val"), str(count))
            row._tr.get_or_add_trPr().append(prop)
    for cell in list(row._tr.tc_lst)[:before]:
        row._tr.remove(cell)
    if after:
        for cell in list(row._tr.tc_lst)[-after:]:
            row._tr.remove(cell)
    for column in range(3):
        if column < before or column >= 3 - after:
            expected_rows[short_row][column] = ""
    assert len(row._tr.tc_lst) == 3 - before - after

    if extension == "txt":
        expected = "\n".join("\t".join(values) for values in expected_rows)
    else:
        lines = ["| " + " | ".join(values) + " |" for values in expected_rows]
        lines.insert(1, "| --- | --- | --- |")
        expected = "\n".join(lines)
    assert convert(document, tmp_path, extension) == expected


@pytest.mark.parametrize("extension", ["txt", "md"])
def test_docx_paragraph_control(tmp_path: Path, extension: str) -> None:
    document = docx.Document()
    document.add_heading("Introduction", level=1)
    document.add_paragraph("Body text.")
    assert convert(document, tmp_path, extension) == (
        "# Introduction\nBody text."
        if extension == "md"
        else "Introduction\nBody text."
    )


def test_docx_markdown_cell_escaping_and_merged_cells(tmp_path: Path) -> None:
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged"
    table.cell(1, 0).text = "pipe|slash\\\nnext <tag>"
    table.cell(1, 1).text = ""
    assert convert(document, tmp_path, "md") == (
        "| Merged | Merged |\n| --- | --- |\n"
        "| pipe\\|slash\\\\<br>next &lt;tag&gt; |  |"
    )


def test_docx_text_multiline_cell(tmp_path: Path) -> None:
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "first\nsecond"
    table.cell(0, 1).text = "other"
    assert convert(document, tmp_path, "txt") == "first\nsecond\tother"


@pytest.mark.parametrize("extension", ["txt", "md"])
def test_docx_consecutive_tables(tmp_path: Path, extension: str) -> None:
    document = docx.Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "First"
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Second"
    document.add_paragraph("After")
    assert convert(document, tmp_path, extension) == (
        "| First |\n| --- |\n\n| Second |\n| --- |\n\nAfter"
        if extension == "md"
        else "First\nSecond\nAfter"
    )


@pytest.mark.parametrize("extension", ["txt", "md"])
def test_docx_vertical_merge_repeats_grid_text(tmp_path: Path, extension: str) -> None:
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "Shared"
    table.cell(0, 1).text = "A"
    table.cell(1, 1).text = "B"
    assert convert(document, tmp_path, extension) == (
        "| Shared | A |\n| --- | --- |\n| Shared | B |"
        if extension == "md"
        else "Shared\tA\nShared\tB"
    )


@pytest.mark.parametrize("extension", ["txt", "md"])
def test_docx_without_row_grid_properties(
    tmp_path: Path, extension: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Table extraction also works with the python-docx 1.1.0 row API."""
    from docx.table import _Row

    document = docx.Document()
    document.add_paragraph("Before")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Cell"
    monkeypatch.delattr(_Row, "grid_cols_before", raising=False)
    monkeypatch.delattr(_Row, "grid_cols_after", raising=False)
    text = convert(document, tmp_path, extension)
    assert text == (
        "Before\n\n| Cell |\n| --- |" if extension == "md" else "Before\nCell"
    )


@pytest.mark.parametrize(
    ("literal", "escaped"),
    [
        ("[draft](https://example.com)", r"\[draft\](https://example.com)"),
        (
            "*required* _literal_ ~~deleted~~",
            r"\*required\* \_literal\_ \~\~deleted\~\~",
        ),
        ("`code` ![image](path)", r"\`code\` !\[image\](path)"),
    ],
)
def test_docx_markdown_literal_cell_text(
    tmp_path: Path, literal: str, escaped: str
) -> None:
    document = docx.Document()
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = literal
    table.cell(1, 0).text = literal + "\nnext"
    assert convert(document, tmp_path, "md") == (
        f"| {escaped} |\n| --- |\n| {escaped}<br>next |"
    )


@pytest.mark.parametrize("extension", ["txt", "md"])
def test_docx_merged_cells_after_omitted_column(tmp_path: Path, extension: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = docx.Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 1).merge(table.cell(2, 2)).text = "Shared"
    for row in table.rows:
        prop = OxmlElement("w:gridBefore")
        prop.set(qn("w:val"), "1")
        row._tr.get_or_add_trPr().append(prop)
        row._tr.remove(row._tr.tc_lst[0])
    assert convert(document, tmp_path, extension) == (
        "|  | Shared | Shared |\n| --- | --- | --- |\n"
        "|  | Shared | Shared |\n|  | Shared | Shared |"
        if extension == "md"
        else "\tShared\tShared\n\tShared\tShared\n\tShared\tShared"
    )
